import sys
import openpyxl

from .database import SessionLocal, Base, engine
from . import models

Base.metadata.create_all(bind=engine)

def clean_text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text if text and text != "-" else None

def clean_number(value):
    if value is None:
        return 0
    try:
        return float(value)
    except (ValueError, TypeError):
        return 0

def import_courses(wb, db):
    ws = wb["التعليم التفاعلي"]
    count = 0
    for row in range(3, ws.max_row + 1):
        name = clean_text(ws.cell(row=row, column=2).value)
        if not name:
            continue
        db.add(models.Course(
            date=clean_text(ws.cell(row=row, column=1).value),
            course_name=name,
            enrolled_count=int(clean_number(ws.cell(row=row, column=3).value)),
            attended_count=int(clean_number(ws.cell(row=row, column=4).value)),
            certificates_issued=int(clean_number(ws.cell(row=row, column=5).value)),
            rating=clean_text(ws.cell(row=row, column=6).value),
            teachers_count=clean_text(ws.cell(row=row, column=7).value),
            reviewed=True,
        ))
        count += 1
    return count

def import_qna(wb, db):
    ws = wb["السؤال والجواب"]
    count = 0
    for row in range(3, ws.max_row + 1):
        question = clean_text(ws.cell(row=row, column=2).value)
        if not question:
            continue
        db.add(models.QnaEntry(
            date=clean_text(ws.cell(row=row, column=1).value),
            question=question,
            status=clean_text(ws.cell(row=row, column=3).value),
            country=clean_text(ws.cell(row=row, column=4).value),
            reviewed=True,
        ))
        count += 1
    return count

def import_projects_and_subscriptions(wb, db):
    ws = wb["المشاريع"]
    projects, subs = 0, 0
    for row in range(3, ws.max_row + 1):
        name = clean_text(ws.cell(row=row, column=2).value)
        if name:
            db.add(models.Project(
                date=clean_text(ws.cell(row=row, column=1).value),
                name=name,
                target_amount=clean_number(ws.cell(row=row, column=3).value),
                raised_amount=clean_number(ws.cell(row=row, column=4).value),
                donors_count=int(clean_number(ws.cell(row=row, column=5).value)),
                completion_percent=clean_text(ws.cell(row=row, column=6).value),
                reviewed=True,
            ))
            projects += 1

        package = clean_text(ws.cell(row=row, column=8).value)
        if package:
            db.add(models.Subscription(
                package=package,
                subscribers_count=int(clean_number(ws.cell(row=row, column=9).value)),
                total_amount=clean_number(ws.cell(row=row, column=10).value),
                reviewed=True,
            ))
            subs += 1
    return projects, subs

def import_communities(wb, db):
    ws = wb["المجتمعات الافتراضية"]
    count = 0
    for row in range(2, ws.max_row + 1):
        name = clean_text(ws.cell(row=row, column=3).value)
        if not name:
            continue
        db.add(models.Community(
            date=clean_text(ws.cell(row=row, column=1).value),
            community_type=clean_text(ws.cell(row=row, column=2).value),
            name=name,
            joined_count=int(clean_number(ws.cell(row=row, column=4).value)),
            posts_count=int(clean_number(ws.cell(row=row, column=5).value)),
            reviewed=True,
        ))
        count += 1
    return count

def import_news(wb, db):
    ws = wb["الأخبار والبيانات"]
    count = 0
    for row in range(2, ws.max_row + 1):
        date_val = clean_text(ws.cell(row=row, column=1).value)
        if not date_val:
            continue
        db.add(models.NewsEntry(
            date=date_val,
            news_count=int(clean_number(ws.cell(row=row, column=2).value)),
            data_count=int(clean_number(ws.cell(row=row, column=3).value)),
            reviewed=True,
        ))
        count += 1
    return count

def import_site_metrics(wb, db, sheet_name, source):
    ws = wb[sheet_name]
    count = 0
    for row in range(1, ws.max_row + 1):
        label = clean_text(ws.cell(row=row, column=1).value)
        if not label:
            continue
        value = clean_text(ws.cell(row=row, column=2).value)
        db.add(models.SiteMetric(source=source, label=label, value=value, reviewed=True))
        count += 1
    return count

def main(file_path: str):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    db = SessionLocal()

    n_courses = import_courses(wb, db)
    n_qna = import_qna(wb, db)
    n_projects, n_subs = import_projects_and_subscriptions(wb, db)
    n_communities = import_communities(wb, db)
    n_news = import_news(wb, db)
    n_rabita = import_site_metrics(wb, db, "بيانات موقع الرابطة", "rabita")
    n_sheikh = import_site_metrics(wb, db, "بيانات موقع معالي الشيخ", "sheikh")
    n_admin = import_site_metrics(wb, db, "الأدمن", "admin_metrics")

    db.commit()

    print("✅ تم الاستيراد بنجاح:")
    print(f"   - التعليم التفاعلي: {n_courses}")
    print(f"   - السؤال والجواب: {n_qna}")
    print(f"   - المشاريع: {n_projects} / الاشتراكات: {n_subs}")
    print(f"   - المجتمعات الافتراضية: {n_communities}")
    print(f"   - الأخبار والبيانات: {n_news}")
    print(f"   - موقع الرابطة: {n_rabita}")
    print(f"   - موقع معالي الشيخ: {n_sheikh}")
    print(f"   - الأدمن: {n_admin}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print('الاستخدام: python -m app.import_all "مسار_ملفك.xlsx"')
        sys.exit(1)
    main(sys.argv[1])
